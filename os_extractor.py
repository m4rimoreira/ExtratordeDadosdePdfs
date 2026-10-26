"""
os_extractor.py
----------------
Automação para:
  1. Ler PDFs de Ordens de Serviço (OS) dentro de uma pasta.
  2. Identificar o início de cada OS pela logo (verde + azul)
     que aparece no topo da página.
  3. Extrair os campos de cada OS (Cliente, Assunto, Cidade, Consultor,
     Obs. do Técnico, Data agendada).
  4. Gerar um relatório .docx com todas as OS's e um ranking de cidades
     com mais Ordens de Serviço.

Instalação das dependências:
    pip install pymupdf python-docx numpy

Uso básico:
    python os_extractor.py "caminho/da/pasta/com/pdfs" -o relatorio_os.docx

Modo de calibração (recomendado na primeira execução):
    python os_extractor.py "caminho/da/pasta" --debug-logo
    -> Gera, em uma pasta "debug_logo/", um recorte da parte superior de
       cada página junto com um .txt informando se a logo foi detectada
       ou não. Use isso para ajustar GREEN_RANGE / BLUE_RANGE abaixo caso
       a detecção esteja errando.
"""

import argparse
import glob
import os
import re
from collections import Counter

import fitz  # PyMuPDF
import numpy as np
from docx import Document

LOGO_TOP_REGION = 0.20          
RENDER_DPI = 150                

GREEN_RANGE = ((20, 100, 20), (150, 220, 150))
BLUE_RANGE = ((10, 60, 120), (120, 170, 255))

MIN_COLOR_PIXELS = 120


def _color_mask(region: np.ndarray, color_range) -> np.ndarray:
    lo = np.array(color_range[0])
    hi = np.array(color_range[1])
    return np.all((region >= lo) & (region <= hi), axis=-1)


def page_has_logo(page: "fitz.Page") -> bool:
    pix = page.get_pixmap(dpi=RENDER_DPI)
    img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)

    top_h = max(1, int(pix.height * LOGO_TOP_REGION))
    region = img[:top_h, :, :3]  # ignora canal alpha, se existir

    green_hits = int(_color_mask(region, GREEN_RANGE).sum())
    blue_hits = int(_color_mask(region, BLUE_RANGE).sum())

    return green_hits >= MIN_COLOR_PIXELS and blue_hits >= MIN_COLOR_PIXELS


def debug_logo_detection(pdf_path: str, out_dir: str = "debug_logo") -> None:
    os.makedirs(out_dir, exist_ok=True)
    doc = fitz.open(pdf_path)
    base_name = os.path.splitext(os.path.basename(pdf_path))[0]

    report_lines = []
    for i, page in enumerate(doc, start=1):
        pix = page.get_pixmap(dpi=RENDER_DPI)
        top_h = max(1, int(pix.height * LOGO_TOP_REGION))
        clip = fitz.Rect(0, 0, page.rect.width, page.rect.height * LOGO_TOP_REGION)
        crop_pix = page.get_pixmap(dpi=RENDER_DPI, clip=clip)
        img_path = os.path.join(out_dir, f"{base_name}_pagina{i:03d}.png")
        crop_pix.save(img_path)

        detectado = page_has_logo(page)
        report_lines.append(f"Página {i}: logo detectada = {detectado}  ({img_path})")

    with open(os.path.join(out_dir, f"{base_name}_relatorio.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(report_lines))

    doc.close()
    print(f"[debug] Recortes e relatório salvos em: {out_dir}/")

LABELS_PDF_TO_REPORT = {
    "Cliente": "Nome do Cliente",
    "Assunto": "Descrição da OS",
    "Cidade": "Localidade",
    "Consultor": "Técnico",
    "Obs. do Técnico": "Observação do Técnico",
    "Data agendada": "Data do Reparo",
}
ORDERED_LABELS = list(LABELS_PDF_TO_REPORT.keys())


def extract_fields(text: str) -> dict:
    labels_alt = "|".join(re.escape(l) for l in ORDERED_LABELS)
    result = {}

    for label in ORDERED_LABELS:
        pattern = rf"{re.escape(label)}\s*:\s*(.*?)(?=(?:{labels_alt})\s*:|\Z)"
        match = re.search(pattern, text, re.DOTALL)
        value = match.group(1).strip() if match else ""
        value = re.sub(r"\s+", " ", value).strip()
        result[label] = value

    return result

def split_into_os_blocks(doc: "fitz.Document") -> list:
    blocks = []
    current_pages_text = []

    for page in doc:
        is_new_os = page_has_logo(page)
        if is_new_os and current_pages_text:
            blocks.append("\n".join(current_pages_text))
            current_pages_text = []
        current_pages_text.append(page.get_text())

    if current_pages_text:
        blocks.append("\n".join(current_pages_text))

    return blocks


def process_pdf(pdf_path: str) -> list:
    doc = fitz.open(pdf_path)
    blocks = split_into_os_blocks(doc)
    doc.close()
    return [extract_fields(block) for block in blocks]


def process_folder(folder: str) -> list:
    all_entries = []
    pdf_files = sorted(glob.glob(os.path.join(folder, "*.pdf")))

    if not pdf_files:
        print(f"Nenhum arquivo .pdf encontrado em: {folder}")
        return all_entries

    for pdf_path in pdf_files:
        print(f"Lendo: {pdf_path}")
        entries = process_pdf(pdf_path)
        print(f"  -> {len(entries)} Ordem(ns) de Serviço encontrada(s)")
        all_entries.extend(entries)

    return all_entries

def build_docx(entries: list, output_path: str) -> None:
    document = Document()
    document.add_heading("Relatório de Ordens de Serviço", level=1)

    for i, entry in enumerate(entries, start=1):
        document.add_heading(f"Ordem de Serviço {i}", level=2)
        for pdf_label, report_label in LABELS_PDF_TO_REPORT.items():
            value = entry.get(pdf_label, "").strip()
            p = document.add_paragraph()
            run_label = p.add_run(f"{report_label}: ")
            run_label.bold = True
            p.add_run(value if value else "(não encontrado)")


    document.add_page_break()
    document.add_heading("Ranking de Cidades com Mais Ordens de Serviço", level=1)

    cidades = [e.get("Cidade", "").strip() for e in entries if e.get("Cidade", "").strip()]
    ranking = Counter(cidades).most_common()

    if ranking:
        table = document.add_table(rows=1, cols=2)
        table.style = "Light Grid Accent 1"
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Cidade"
        hdr_cells[1].text = "Qtd. de Ordens de Serviço"

        for cidade, qtd in ranking:
            row_cells = table.add_row().cells
            row_cells[0].text = cidade
            row_cells[1].text = str(qtd)
    else:
        document.add_paragraph("Nenhuma cidade foi identificada nas Ordens de Serviço.")

    document.save(output_path)


def main():
    parser = argparse.ArgumentParser(
        description="Extrai Ordens de Serviço de PDFs e gera um relatório .docx"
    )
    parser.add_argument("pasta_pdfs", help="Pasta contendo os arquivos PDF das Ordens de Serviço")
    parser.add_argument("-o", "--saida", default="relatorio_os.docx", help="Caminho do .docx de saída")
    parser.add_argument(
        "--debug-logo",
        action="store_true",
        help="Gera recortes de depuração da faixa superior de cada página, sem gerar o relatório",
    )
    args = parser.parse_args()

    if args.debug_logo:
        pdf_files = sorted(glob.glob(os.path.join(args.pasta_pdfs, "*.pdf")))
        if not pdf_files:
            print(f"Nenhum .pdf encontrado em: {args.pasta_pdfs}")
            return
        for pdf_path in pdf_files:
            debug_logo_detection(pdf_path)
        return

    entries = process_folder(args.pasta_pdfs)
    if not entries:
        print("Nenhuma Ordem de Serviço foi extraída. Rode com --debug-logo para calibrar a detecção do logo.")
        return

    build_docx(entries, args.saida)
    print(f"\n{len(entries)} Ordem(ns) de Serviço processada(s).")
    print(f"Relatório salvo em: {args.saida}")


if __name__ == "__main__":
    main()
